import openpyxl
wb=openpyxl.load_workbook("parmita_atm.xlsx")
sheet=wb.active
while True:
    def menu():
        print("==Welcome to ATM system==")
        print("1.Create account ")
        print("2.Deposite money ")
        print("3.withdraw money ")
        print("4.View balance ")
        print("5.Exit ")
    menu()
    select=int(input("select the option "))
    if select==1:
        def create_account():
            phone_no=int(input("Enter your phone number "))
            name=input("Enter your name ")
            deposite_balance=int(input("Enter the deposite balance "))
            sheet.append([name,phone_no,deposite_balance])
            wb.save("parmita_atm.xlsx")
            print("account is successfully created")
        create_account()
    elif select==2:
        def deposite_money():
            phone_no=input("Enter your phone number ")
            add_money=int(input("Enter the account to deposite "))
            for row in sheet.iter_rows(min_row=2):
                if str(row[1].value)==phone_no:
                    row[2].value=row[2].value+add_money
                    wb.save("parmita_atm.xlsx")
                    print("Amount deposited successfully")
                    print("New balance:",row[2].value)
            return
            print("Invalid account")
        deposite_money()
    elif select==3:
        def withdraw():
              phone_no=input("Enter your phone number ")
              withdraw_money=int(input("Enter amount to withdraw "))
              for row in sheet.iter_rows(min_row=2):
                  if str(row[1].value)==phone_no:
                      balance=row[2].value
                      if withdraw_money<=0:
                          print("Enter a valid amount")
                      elif withdraw_money<=balance:
                          row[2].value=balance-withdraw_money
                          wb.save("parmita_atm.xlsx")
                          print("withdraw successful")
                          print("Remaining balance:",row[2].value)
                      else:
                         print("Insufficient balance")
                      return
                      print("Invalid account")
        withdraw()
    elif select==4:
        def view_balance():
            phone_no=input("Enter your phone number ")
            found=False
            for row in sheet.iter_rows(min_row=2,values_only=True):
                if str(row[1])==phone_no:
                    print("Your balance is: ",row[2])
                    found=True
                    break
            if not found:
                print("Invalid account!")
        view_balance()
    elif select==5:
        print("Thankyou for using ATM system")
        print("Exit")
        break
    else:
        print("Invalid choice")
        
            


        
        
        
        



